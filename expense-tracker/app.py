import os
import json
import sqlite3
from datetime import datetime
from pathlib import Path
import base64
from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, send_file, Response
import anthropic
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import io

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'hikmah-expense-tracker-secret')

UPLOAD_FOLDER = Path(os.environ.get('UPLOAD_FOLDER', 'uploads'))
UPLOAD_FOLDER.mkdir(exist_ok=True)

DB_PATH = os.environ.get('DB_PATH', 'expenses.db')

CATEGORIES = [
    'Electricity', 'Gas', 'Water', 'Internet', 'Phone',
    'Club Membership', 'Salary', 'Petrol', 'Groceries', 'Other'
]

ALLOWED_EXTENSIONS = {'.jpg', '.jpeg', '.png', '.webp', '.gif'}

CATEGORY_COLORS = {
    'Electricity': ('#fef9c3', '#854d0e'),
    'Gas':         ('#fef3c7', '#92400e'),
    'Water':       ('#dbeafe', '#1e40af'),
    'Internet':    ('#f0fdf4', '#166534'),
    'Phone':       ('#f0fdf4', '#166534'),
    'Club Membership': ('#fdf2f8', '#86198f'),
    'Salary':      ('#ecfdf5', '#065f46'),
    'Petrol':      ('#fff7ed', '#9a3412'),
    'Groceries':   ('#f0fdfa', '#134e4a'),
    'Other':       ('#f3f4f6', '#374151'),
}


def init_db():
    conn = sqlite3.connect(DB_PATH)
    conn.execute('''
        CREATE TABLE IF NOT EXISTS bills (
            id           INTEGER PRIMARY KEY AUTOINCREMENT,
            filename     TEXT,
            upload_date  TEXT NOT NULL,
            bill_date    TEXT,
            vendor       TEXT,
            category     TEXT NOT NULL DEFAULT 'Other',
            amount       REAL NOT NULL DEFAULT 0,
            currency     TEXT NOT NULL DEFAULT 'PKR',
            description  TEXT,
            notes        TEXT
        )
    ''')
    conn.commit()
    conn.close()


init_db()


def require_auth(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        password = os.environ.get('DASHBOARD_PASSWORD', 'hikmah2026')
        auth = request.authorization
        if not auth or auth.password != password:
            return Response(
                'Login required.',
                401,
                {'WWW-Authenticate': 'Basic realm="Expense Dashboard"'}
            )
        return f(*args, **kwargs)
    return decorated


def analyze_bill(image_path: Path) -> dict:
    api_key = os.environ.get('ANTHROPIC_API_KEY')
    if not api_key:
        return {
            'vendor': 'Unknown',
            'category': 'Other',
            'amount': 0.0,
            'currency': 'PKR',
            'bill_date': '',
            'description': 'Set ANTHROPIC_API_KEY to enable automatic bill reading'
        }

    client = anthropic.Anthropic(api_key=api_key)

    with open(image_path, 'rb') as f:
        image_data = base64.b64encode(f.read()).decode('utf-8')

    media_type_map = {
        '.jpg': 'image/jpeg', '.jpeg': 'image/jpeg',
        '.png': 'image/png', '.webp': 'image/webp', '.gif': 'image/gif'
    }
    media_type = media_type_map.get(image_path.suffix.lower(), 'image/jpeg')

    try:
        response = client.messages.create(
            model='claude-haiku-4-5-20251001',
            max_tokens=600,
            messages=[{
                'role': 'user',
                'content': [
                    {
                        'type': 'image',
                        'source': {
                            'type': 'base64',
                            'media_type': media_type,
                            'data': image_data
                        }
                    },
                    {
                        'type': 'text',
                        'text': (
                            'Analyze this bill or receipt image from Pakistan.\n'
                            'Extract the key information and return ONLY a JSON object '
                            '(no markdown, no explanation).\n\n'
                            f'Categories to choose from: {", ".join(CATEGORIES)}\n\n'
                            'Return exactly this structure:\n'
                            '{\n'
                            '  "vendor": "company or utility name",\n'
                            '  "category": "one category from the list above",\n'
                            '  "amount": 0.00,\n'
                            '  "currency": "PKR",\n'
                            '  "bill_date": "YYYY-MM-DD or empty string if not clear",\n'
                            '  "description": "short description e.g. May 2026 electricity bill"\n'
                            '}\n\n'
                            'amount must be a number (the total amount due on the bill). '
                            'Use 0 if you cannot read it clearly.'
                        )
                    }
                ]
            }]
        )

        text = response.content[0].text.strip()

        # Strip markdown code fences if present
        if '```' in text:
            for part in text.split('```'):
                part = part.strip().lstrip('json').strip()
                try:
                    return json.loads(part)
                except (json.JSONDecodeError, ValueError):
                    continue

        return json.loads(text)

    except Exception as e:
        return {
            'vendor': 'Unknown',
            'category': 'Other',
            'amount': 0.0,
            'currency': 'PKR',
            'bill_date': '',
            'description': f'Could not read bill automatically ({str(e)[:80]})'
        }


@app.route('/')
def upload_page():
    return render_template('upload.html')


@app.route('/upload', methods=['POST'])
def upload_bill():
    if 'bill' not in request.files or not request.files['bill'].filename:
        return render_template('upload.html', error='No photo selected. Please choose a photo of the bill.')

    file = request.files['bill']
    ext = Path(file.filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        return render_template('upload.html', error='Please upload a JPG or PNG photo.')

    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f'{timestamp}{ext}'
    save_path = UPLOAD_FOLDER / filename
    file.save(save_path)

    data = analyze_bill(save_path)

    conn = sqlite3.connect(DB_PATH)
    cursor = conn.execute(
        '''INSERT INTO bills
           (filename, upload_date, bill_date, vendor, category, amount, currency, description)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?)''',
        (
            filename,
            datetime.now().strftime('%Y-%m-%d'),
            data.get('bill_date') or '',
            data.get('vendor') or 'Unknown',
            data.get('category') or 'Other',
            float(data.get('amount') or 0),
            data.get('currency') or 'PKR',
            data.get('description') or '',
        )
    )
    conn.commit()
    conn.close()

    return render_template('upload.html', success=True, data=data)


@app.route('/dashboard')
@require_auth
def dashboard():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    bills = conn.execute(
        'SELECT * FROM bills ORDER BY upload_date DESC, id DESC'
    ).fetchall()
    conn.close()

    monthly = {}
    for bill in bills:
        key = bill['upload_date'][:7] if bill['upload_date'] else 'Unknown'
        if key not in monthly:
            try:
                label = datetime.strptime(key, '%Y-%m').strftime('%B %Y')
            except ValueError:
                label = key
            monthly[key] = {'label': label, 'bills': [], 'total': 0.0}
        monthly[key]['bills'].append(dict(bill))
        monthly[key]['total'] += float(bill['amount'] or 0)

    return render_template(
        'dashboard.html',
        monthly=monthly,
        categories=CATEGORIES,
        category_colors=CATEGORY_COLORS
    )


@app.route('/image/<filename>')
@require_auth
def serve_image(filename):
    path = UPLOAD_FOLDER / filename
    if not path.exists():
        return 'Image not found', 404
    return send_file(path)


@app.route('/edit/<int:bill_id>', methods=['GET', 'POST'])
@require_auth
def edit_bill(bill_id):
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row

    if request.method == 'POST':
        conn.execute(
            '''UPDATE bills
               SET vendor=?, category=?, amount=?, bill_date=?, description=?, notes=?
               WHERE id=?''',
            (
                request.form.get('vendor', ''),
                request.form.get('category', 'Other'),
                float(request.form.get('amount', 0) or 0),
                request.form.get('bill_date', ''),
                request.form.get('description', ''),
                request.form.get('notes', ''),
                bill_id
            )
        )
        conn.commit()
        conn.close()
        return redirect(url_for('dashboard'))

    bill = conn.execute('SELECT * FROM bills WHERE id=?', (bill_id,)).fetchone()
    conn.close()
    if not bill:
        return redirect(url_for('dashboard'))
    return render_template('edit.html', bill=dict(bill), categories=CATEGORIES)


@app.route('/delete/<int:bill_id>', methods=['POST'])
@require_auth
def delete_bill(bill_id):
    conn = sqlite3.connect(DB_PATH)
    row = conn.execute('SELECT filename FROM bills WHERE id=?', (bill_id,)).fetchone()
    if row:
        img = UPLOAD_FOLDER / row[0]
        if img.exists():
            img.unlink()
        conn.execute('DELETE FROM bills WHERE id=?', (bill_id,))
        conn.commit()
    conn.close()
    return redirect(url_for('dashboard'))


@app.route('/export')
@require_auth
def export_excel():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    bills = conn.execute(
        'SELECT * FROM bills ORDER BY upload_date, category'
    ).fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Expenses'

    header_fill = PatternFill(start_color='1A3C5E', end_color='1A3C5E', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=11)
    alt_fill = PatternFill(start_color='EEF4FB', end_color='EEF4FB', fill_type='solid')
    thin = Side(style='thin', color='CCCCCC')
    row_border = Border(bottom=thin)

    headers = ['Upload Date', 'Bill Date', 'Vendor', 'Category', 'Description', 'Amount (PKR)', 'Notes']
    col_widths = [14, 12, 22, 18, 38, 16, 25]

    for col, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[1].height = 26

    for row_num, bill in enumerate(bills, 2):
        fill = alt_fill if row_num % 2 == 0 else None
        row_data = [
            bill['upload_date'],
            bill['bill_date'] or '',
            bill['vendor'],
            bill['category'],
            bill['description'] or '',
            float(bill['amount'] or 0),
            bill['notes'] or '',
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col, value=val)
            if fill:
                cell.fill = fill
            cell.border = row_border
            if col == 6:
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal='right')

    if bills:
        total_row = len(bills) + 2
        label = ws.cell(row=total_row, column=5, value='TOTAL')
        label.font = Font(bold=True)
        total = ws.cell(row=total_row, column=6, value=f'=SUM(F2:F{total_row - 1})')
        total.font = Font(bold=True)
        total.number_format = '#,##0.00'
        total.alignment = Alignment(horizontal='right')

    ws.freeze_panes = 'A2'

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    month = datetime.now().strftime('%B_%Y')
    return send_file(
        output,
        download_name=f'expenses_{month}.xlsx',
        as_attachment=True,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )


if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=False)
